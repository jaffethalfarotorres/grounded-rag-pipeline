import pytest

from grounded_rag.loaders import UnsupportedFormat, load_directory, load_path, slugify
from tests.pdf_fixture import minimal_pdf

MARKDOWN = """---
category: billing
owner: finance
---
# Refund policy

Refunds are issued within 14 days of the request.
"""


def test_markdown_frontmatter_becomes_metadata_and_heading_becomes_title(tmp_path):
    path = tmp_path / "refunds.md"
    path.write_text(MARKDOWN, encoding="utf-8")
    doc = load_path(path)
    assert doc.doc_id == "refunds"
    assert doc.title == "Refund policy"
    assert doc.metadata == {"category": "billing", "owner": "finance"}
    assert "category:" not in doc.text
    assert "within 14 days" in doc.text


def test_plain_text_title_comes_from_the_file_name(tmp_path):
    path = tmp_path / "on_call_rotation.txt"
    path.write_text("Engineers rotate weekly.", encoding="utf-8")
    doc = load_path(path)
    assert doc.title == "On call rotation"
    assert doc.text == "Engineers rotate weekly."


def test_docx_paragraphs_are_extracted(tmp_path):
    from docx import Document as DocxDocument

    path = tmp_path / "security.docx"
    word = DocxDocument()
    word.add_paragraph("Data is encrypted at rest with AES-256.")
    word.add_paragraph("Keys rotate every 90 days.")
    word.save(path)
    doc = load_path(path)
    assert "AES-256" in doc.text
    assert "\n\n" in doc.text


def test_pdf_text_is_extracted(tmp_path):
    path = tmp_path / "sla.pdf"
    path.write_bytes(minimal_pdf("Monthly uptime target is 99.9 percent."))
    assert "99.9 percent" in load_path(path).text


def test_unsupported_format_is_rejected(tmp_path):
    path = tmp_path / "photo.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(UnsupportedFormat, match="png"):
        load_path(path)


@pytest.mark.parametrize(
    ("name", "slug"),
    [
        ("Refund Policy (v2)", "refund-policy-v2"),
        ("SLA", "sla"),
        ("data_retention", "data_retention"),
    ],
)
def test_slugify_makes_valid_document_ids(name, slug):
    assert slugify(name) == slug


def test_directory_loading_is_sorted_and_reports_skipped_files(tmp_path):
    (tmp_path / "b.md").write_text("# B\n\nbeta", encoding="utf-8")
    (tmp_path / "a.txt").write_text("alpha", encoding="utf-8")
    (tmp_path / "notes.png").write_bytes(b"\x89PNG")
    documents, skipped = load_directory(tmp_path)
    assert [doc.doc_id for doc in documents] == ["a", "b"]
    assert [path.name for path in skipped] == ["notes.png"]
